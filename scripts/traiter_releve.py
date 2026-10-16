#!/usr/bin/env python3
"""Traite les réponses au relevé collaboratif de la pleine terre (releve-pleine-terre.html).

Usage : python scripts/traiter_releve.py [dossier]   (par défaut : reponses/)

Sources : les fichiers .json du dossier, et la collecte en ligne (feuille Google, voir
collecte/LISEZMOI.md) si les variables RELEVE_URL et RELEVE_CLE sont définies.

Placez dans le dossier les fichiers .json reçus des participants. Si une même personne a
envoyé plusieurs fichiers, seul le plus récent compte (chaque envoi contient toutes ses réponses).

Pour chaque terrain, on retient la médiane des parts de pleine terre estimées par les
participants (avant et après), l'état du projet le plus souvent cité, et la médiane des
surfaces dessinées, limitées au terrain.

Produit :
- reponses/synthese-releve.xlsx (ou .csv sans openpyxl) : détail par terrain, avec les
  commentaires et la liste des participants. À garder pour vous : il contient les prénoms.
- data/releve_pleine_terre.js et .json : résultats anonymes, que le site peut afficher.
"""

import csv
import datetime as dt
import json
import math
import os
import pathlib
import re
import statistics
import sys
import urllib.parse
import urllib.request
from collections import Counter

RACINE = pathlib.Path(__file__).resolve().parent.parent
DATA = RACINE / "data"
PAS_GRILLE = 0.5  # m : finesse du calcul des surfaces dessinées à l'intérieur du terrain
SEUIL_DESACCORD = 25  # points de pourcentage d'écart entre participants


# --------------------------------------------------------------------------- Géométrie

def projeter(lat0, lon0):
    kx = 111320 * math.cos(math.radians(lat0))
    return lambda lat, lon: ((lon - lon0) * kx, (lat - lat0) * 110574)


def dans(x, y, anneau):
    dedans = False
    j = len(anneau) - 1
    for i in range(len(anneau)):
        xi, yi = anneau[i]
        xj, yj = anneau[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            dedans = not dedans
        j = i
    return dedans


def dans_terrain(x, y, polys):
    for poly in polys:
        if dans(x, y, poly[0]) and not any(dans(x, y, trou) for trou in poly[1:]):
            return True
    return False


def surface_dans_terrain(points, polys_terrain, proj):
    """Surface (m²) de la zone dessinée qui tombe à l'intérieur du terrain, par comptage sur une grille."""
    zone = [proj(lat, lon) for lat, lon in points]
    xs, ys = [p[0] for p in zone], [p[1] for p in zone]
    n = 0
    y = min(ys) + PAS_GRILLE / 2
    while y < max(ys):
        x = min(xs) + PAS_GRILLE / 2
        while x < max(xs):
            if dans(x, y, zone) and dans_terrain(x, y, polys_terrain):
                n += 1
            x += PAS_GRILLE
        y += PAS_GRILLE
    return n * PAS_GRILLE ** 2


def charger_permis():
    """Contours des parcelles, et état des autorisations selon la base Sitadel."""
    paquet = json.loads((DATA / "permis.json").read_text(encoding="utf-8"))
    contours = {}
    for p in paquet["permis"]:
        for c in p.get("contours") or []:
            contours[c["parcelle"]] = c
    return contours, {p["id"]: p["etat"] for p in paquet["permis"]}


def etat_sitadel(ids, etats):
    """État du projet d'après Sitadel, quand aucun participant ne l'a indiqué."""
    connus = [etats.get(i) for i in ids]
    if "termine" in connus:
        return "termine"
    if "commence" in connus:
        return "chantier"
    return None


# --------------------------------------------------------------------------- Lecture des réponses

def lire_reponses(dossier):
    par_participant = {}
    for f in sorted(dossier.glob("*.json")):
        try:
            d = json.loads(f.read_text(encoding="utf-8"))
        except (OSError, ValueError) as e:
            print(f"  ignoré : {f.name} ({e})")
            continue
        if d.get("format") != "releve-pleine-terre-gentilly":
            print(f"  ignoré : {f.name} (ce n'est pas un fichier du relevé)")
            continue
        nom = (d.get("participant") or f.stem).strip()
        cle = nom.lower()
        if cle not in par_participant or d.get("exporte_le", "") > par_participant[cle].get("exporte_le", ""):
            par_participant[cle] = d
    return list(par_participant.values())


def mediane(valeurs):
    v = [x for x in valeurs if x is not None]
    return statistics.median(v) if v else None


def synthese(fichiers, contours, etats_permis):
    terrains = {}
    for d in fichiers:
        for r in d.get("reponses", []):
            t = terrains.setdefault(r["terrain"], {"terrain": r["terrain"], "adresse": r.get("adresse"), "parcelles": r.get("parcelles", []),
                                                   "permis": r.get("permis", []), "surface_terrain_m2": r.get("surface_terrain_m2"), "reponses": []})
            t["reponses"].append((d.get("participant", "?"), r))

    lignes = []
    for t in terrains.values():
        geoms = [contours[k] for k in t["parcelles"] if k in contours]
        polys = []
        proj = None
        if geoms:
            lat0 = geoms[0]["coordinates"][0][0][1] if geoms[0]["type"] == "Polygon" else geoms[0]["coordinates"][0][0][0][1]
            lon0 = geoms[0]["coordinates"][0][0][0] if geoms[0]["type"] == "Polygon" else geoms[0]["coordinates"][0][0][0][0]
            proj = projeter(lat0, lon0)
            for g in geoms:
                for poly in (g["coordinates"] if g["type"] == "MultiPolygon" else [g["coordinates"]]):
                    polys.append([[proj(lat, lon) for lon, lat in anneau] for anneau in poly])
        surface = t["surface_terrain_m2"] or 0
        avant = [r.get("part_avant") for _, r in t["reponses"]]
        apres = [r.get("part_apres") for _, r in t["reponses"]]
        etats = Counter(r.get("etat_projet") for _, r in t["reponses"] if r.get("etat_projet") not in (None, "inconnu"))
        disparues, creees = [], []
        for _, r in t["reponses"]:
            zones = r.get("zones") or []
            if not zones or not polys:
                continue
            disparues.append(sum(surface_dans_terrain(z["points"], polys, proj) for z in zones if z["type"] == "disparue"))
            creees.append(sum(surface_dans_terrain(z["points"], polys, proj) for z in zones if z["type"] == "creee"))
        densites = Counter(r.get("densite") for _, r in t["reponses"] if r.get("densite") not in (None, "inconnu"))
        notes = [r.get("appreciation") for _, r in t["reponses"]]
        m_avant, m_apres = mediane(avant), mediane(apres)
        ecarts = [max(v) - min(v) for v in ([x for x in avant if x is not None], [x for x in apres if x is not None]) if len(v) > 1]
        lignes.append({
            "terrain": t["terrain"],
            "adresse": t["adresse"],
            "parcelles": ", ".join(t["parcelles"]),
            "permis": ", ".join(t["permis"]),
            "surface_terrain_m2": surface,
            "nb_reponses": len(t["reponses"]),
            "part_avant_pct": m_avant,
            "part_apres_pct": m_apres,
            "pleine_terre_avant_m2": round(m_avant / 100 * surface) if m_avant is not None else None,
            "pleine_terre_apres_m2": round(m_apres / 100 * surface) if m_apres is not None else None,
            "evolution_m2": round((m_apres - m_avant) / 100 * surface) if None not in (m_avant, m_apres) else None,
            "desaccord": bool(ecarts and max(ecarts) >= SEUIL_DESACCORD),
            "etat_projet": etats.most_common(1)[0][0] if etats else etat_sitadel(t["permis"], etats_permis),
            "etat_source": "participants" if etats else "Sitadel",
            "densite": densites.most_common(1)[0][0] if densites else None,
            "nb_avis_densite": sum(densites.values()),
            "appreciation": mediane(notes),
            "nb_avis_appreciation": len([n for n in notes if n is not None]),
            "nb_dessins": len(disparues),
            "dessin_disparue_m2": round(statistics.median(disparues)) if disparues else None,
            "dessin_creee_m2": round(statistics.median(creees)) if creees else None,
            "commentaires": " | ".join(f"{nom} : {r['commentaire'].strip()}" for nom, r in t["reponses"] if (r.get("commentaire") or "").strip()),
        })
    lignes.sort(key=lambda l: (l["evolution_m2"] is None, l["evolution_m2"] or 0))
    return lignes


# --------------------------------------------------------------------------- Écriture

COLONNES = [
    ("adresse", "Adresse"), ("parcelles", "Parcelles"), ("permis", "Autorisations"), ("surface_terrain_m2", "Surface du terrain (m²)"),
    ("nb_reponses", "Réponses"), ("part_avant_pct", "Pleine terre avant (%)"), ("part_apres_pct", "Pleine terre après (%)"),
    ("pleine_terre_avant_m2", "Pleine terre avant (m²)"), ("pleine_terre_apres_m2", "Pleine terre après (m²)"),
    ("evolution_m2", "Évolution (m²)"), ("desaccord", "Avis divergents"), ("etat_projet", "Projet sur la photo après"),
    ("etat_source", "État d'après"),
    ("densite", "Avis : densité"), ("nb_avis_densite", "Avis densité (nb)"), ("appreciation", "Avis : note sur 5 (médiane)"),
    ("nb_avis_appreciation", "Notes (nb)"),
    ("nb_dessins", "Dessins"), ("dessin_disparue_m2", "Zones dessinées disparues (m²)"), ("dessin_creee_m2", "Zones dessinées créées (m²)"),
    ("commentaires", "Commentaires"),
]
ETATS = {"termine": "Terminé", "chantier": "En chantier", "pas_commence": "Pas commencé"}
DENSITES = {"beaucoup_plus": "Beaucoup plus dense", "plus": "Un peu plus dense", "pareil": "Aussi dense", "moins": "Moins dense"}


def valeur(cle, v):
    if cle == "desaccord":
        return "oui" if v else ""
    if cle == "etat_projet":
        return ETATS.get(v, "")
    if cle == "densite":
        return DENSITES.get(v, "")
    return v


def ecrire_tableur(lignes, fichiers, dossier):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        chemin = dossier / "synthese-releve.csv"
        with open(chemin, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow([titre for _, titre in COLONNES])
            for l in lignes:
                w.writerow([valeur(c, l[c]) if l[c] is not None else "" for c, _ in COLONNES])
        return chemin
    wb = Workbook()
    ws = wb.active
    ws.title = "Terrains"
    for col, (_, titre) in enumerate(COLONNES, 1):
        c = ws.cell(row=1, column=col, value=titre)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="3155A6")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, l in enumerate(lignes, 2):
        for col, (cle, _) in enumerate(COLONNES, 1):
            c = ws.cell(row=i, column=col, value=valeur(cle, l[cle]))
            c.font = Font(name="Arial", size=10, color="E30613" if cle == "evolution_m2" and (l[cle] or 0) < 0 else "000000")
            if cle.endswith("_m2"):
                c.number_format = "#,##0"
    fin = len(lignes) + 1
    total = fin + 2
    ws.cell(row=total, column=1, value="Total des terrains relevés").font = Font(name="Arial", bold=True)
    for cle in ("surface_terrain_m2", "pleine_terre_avant_m2", "pleine_terre_apres_m2", "evolution_m2", "dessin_disparue_m2", "dessin_creee_m2"):
        col = [c for c, _ in COLONNES].index(cle) + 1
        lettre = ws.cell(row=1, column=col).column_letter
        c = ws.cell(row=total, column=col, value=f"=SUM({lettre}2:{lettre}{fin})")
        c.font = Font(name="Arial", bold=True)
        c.number_format = "#,##0"
    ws.cell(row=total + 1, column=1, value="Pertes seules (évolutions négatives)").font = Font(name="Arial", bold=True)
    col = [c for c, _ in COLONNES].index("evolution_m2") + 1
    lettre = ws.cell(row=1, column=col).column_letter
    c = ws.cell(row=total + 1, column=col, value=f'=SUMIF({lettre}2:{lettre}{fin},"<0")')
    c.font = Font(name="Arial", bold=True, color="E30613")
    c.number_format = "#,##0"
    for col, largeur in enumerate([30, 14, 30, 12, 9, 11, 11, 12, 12, 12, 10, 14, 12, 18, 10, 12, 9, 9, 13, 13, 60], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = largeur
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLONNES)).column_letter}{fin}"
    ws.cell(row=total + 3, column=1, value="Valeurs par terrain : médiane des réponses. Surfaces en m² = part estimée × surface du terrain. "
            "Zones dessinées : médiane des participants ayant dessiné, limitée au terrain.").font = Font(name="Arial", italic=True, size=9)

    wp = wb.create_sheet("Participants")
    wp.append(["Participant", "Terrains relevés", "Fichier envoyé le"])
    for d in fichiers:
        wp.append([d.get("participant"), len(d.get("reponses", [])), (d.get("exporte_le") or "")[:16].replace("T", " ")])
    for cell in wp[1]:
        cell.font = Font(name="Arial", bold=True)
    wb.calculation.fullCalcOnLoad = True  # Excel calcule les totaux à l'ouverture
    chemin = dossier / "synthese-releve.xlsx"
    wb.save(chemin)
    return chemin


def min_reponses():
    """Seuil de réponses pour publier un terrain, lu dans data/config.js (3 par défaut)."""
    try:
        m = re.search(r"min_reponses:\s*(\d+)", (DATA / "config.js").read_text(encoding="utf-8"))
        return int(m.group(1)) if m else 3
    except OSError:
        return 3


def ecrire_public(lignes, fichiers, seuil):
    """Résultats anonymes pour le site : ni pseudos, ni commentaires des participants.
    Seuls les terrains ayant au moins `seuil` réponses sont publiés, pour limiter les réponses fantaisistes."""
    publies = [l for l in lignes if l["nb_reponses"] >= seuil]
    construits = [l for l in publies if l["etat_projet"] in ("termine", "chantier") and l["evolution_m2"] is not None]
    paquet = {
        "genere_le": dt.date.today().isoformat(),
        "participants": len(fichiers),
        "min_reponses": seuil,
        "terrains_releves": len(publies),
        "terrains_en_attente": len(lignes) - len(publies),
        "terrains_construits": len(construits),
        "pertes_m2": -sum(l["evolution_m2"] for l in construits if l["evolution_m2"] < 0),
        "gains_m2": sum(l["evolution_m2"] for l in construits if l["evolution_m2"] > 0),
        "terrains": {l["terrain"]: {k: l[k] for k in ("nb_reponses", "part_avant_pct", "part_apres_pct", "pleine_terre_avant_m2",
                                                      "pleine_terre_apres_m2", "evolution_m2", "desaccord", "etat_projet",
                                                      "dessin_disparue_m2", "dessin_creee_m2", "densite", "nb_avis_densite",
                                                      "appreciation", "nb_avis_appreciation")} for l in publies},
    }
    (DATA / "releve_pleine_terre.json").write_text(json.dumps(paquet, ensure_ascii=False, indent=1), encoding="utf-8")
    (DATA / "releve_pleine_terre.js").write_text(
        "/* Fichier généré par scripts/traiter_releve.py — ne pas modifier à la main. */\n"
        "window.RELEVE_PLEINE_TERRE = " + json.dumps(paquet, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    return paquet


def lire_collecte(url, cle):
    """Réponses et commentaires validés, depuis la feuille Google (collecte/Code.gs)."""
    with urllib.request.urlopen(f"{url}?{urllib.parse.urlencode({'cle': cle})}", timeout=120) as rep:
        d = json.load(rep)
    if not d.get("ok"):
        sys.exit(f"Collecte refusée : {d.get('erreur')}")
    par_participant = {}
    for ligne in d.get("releve", []):
        try:
            r = json.loads(ligne["reponse_json"])
        except (TypeError, ValueError):
            continue
        pid = str(ligne.get("participant_id") or "?")
        p = par_participant.setdefault(pid, {"format": "releve-pleine-terre-gentilly", "participant": ligne.get("pseudo") or f"anonyme-{pid[:6]}",
                                              "exporte_le": "", "par_terrain": {}})
        p["exporte_le"] = max(p["exporte_le"], ligne.get("recu_le") or "")
        if ligne.get("pseudo"):
            p["participant"] = ligne["pseudo"]
        p["par_terrain"][r.get("terrain")] = r  # les lignes arrivent dans l'ordre : la dernière réponse l'emporte
    fichiers = [{**{k: v for k, v in p.items() if k != "par_terrain"}, "reponses": list(p["par_terrain"].values())} for p in par_participant.values()]
    commentaires = {}
    for c in d.get("commentaires", []):
        commentaires.setdefault(c["permis"], []).append({"pseudo": c.get("pseudo") or "", "texte": c["texte"], "date": (c.get("date") or "")[:10]})
    return fichiers, commentaires


def ecrire_commentaires(commentaires):
    (DATA / "commentaires.js").write_text(
        "/* Fichier généré par scripts/traiter_releve.py — commentaires validés par la modération. */\n"
        "window.COMMENTAIRES = " + json.dumps(commentaires, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    dossier = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else RACINE / "reponses"
    fichiers = lire_reponses(dossier) if dossier.is_dir() else []
    url, cle = os.environ.get("RELEVE_URL"), os.environ.get("RELEVE_CLE")
    if url and cle:
        en_ligne, commentaires = lire_collecte(url, cle)
        fichiers += en_ligne
        ecrire_commentaires(commentaires)
        print(f"Collecte en ligne : {len(en_ligne)} participant(s), {sum(len(v) for v in commentaires.values())} commentaire(s) publié(s).")
    if not fichiers:
        if url and cle:
            ecrire_public([], [], min_reponses())  # résultats vides, mais datés du jour
            print("Aucune réponse pour le moment dans la collecte en ligne.")
            return
        sys.exit(f"Aucune réponse : ni fichier dans {dossier}, ni collecte en ligne (RELEVE_URL et RELEVE_CLE).")
    contours, etats_permis = charger_permis()
    lignes = synthese(fichiers, contours, etats_permis)
    seuil = min_reponses()
    pub = ecrire_public(lignes, fichiers, seuil)
    print(f"{len(fichiers)} participant(s), {len(lignes)} terrain(s) relevé(s), dont {pub['terrains_releves']} publiés (au moins {seuil} réponses).")
    print(f"Terrains publiés construits : {pub['terrains_construits']} ; pleine terre perdue : {pub['pertes_m2']} m², gagnée : {pub['gains_m2']} m².")
    if dossier.is_dir():
        print(f"Détail (tous les terrains, avec les pseudos) : {ecrire_tableur(lignes, fichiers, dossier)}")
    print("Résultats anonymes pour le site : data/releve_pleine_terre.js")


if __name__ == "__main__":
    main()
